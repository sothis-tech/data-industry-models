# rag-manager / backend / main.py
import json
import os
import re
import shutil
from collections import defaultdict
from pathlib import Path
from typing import Optional

import chromadb, httpx
from fastapi import FastAPI, File, HTTPException, UploadFile, Form
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel

# ── Config ────────────────────────────────────────────────────────────────────
CHROMA_HOST     = os.getenv("CHROMA_HOST",       "localhost")
CHROMA_PORT     = int(os.getenv("CHROMA_PORT",   "8200"))
COLLECTION_NAME = os.getenv("CHROMA_COLLECTION", "manuales_tecnicos")
DOCS_BASE_PATH  = Path(os.getenv("DOCS_PATH",    "./docs"))
CHROMA_DATABASE = "default_database"
TENANTS_FILE    = DOCS_BASE_PATH / ".tenants.json"

# Extensiones soportadas
SUPPORTED_EXTS = {".pdf", ".docx", ".txt", ".md"}

# ══════════════════════════════════════════════════════════════════════════════
# MODELO DE EMBEDDINGS
# ══════════════════════════════════════════════════════════════════════════════
# IMPORTANTE: este modelo DEBE ser exactamente el mismo que el del rag_server.py
# (incluidos los prefijos query/passage). Si no, los embeddings de query y los
# de los documentos viven en espacios distintos y el retrieval no funciona.
#
# intfloat/multilingual-e5-small:
#   - Multilingüe (español/inglés/50+ idiomas)
#   - 384 dimensiones, ventana de 512 tokens
#   - Requiere prefijos: "query: " para preguntas, "passage: " para documentos
EMBED_MODEL_NAME = "intfloat/multilingual-e5-small"

try:
    from llama_index.core import Settings, StorageContext, VectorStoreIndex, Document
    from llama_index.core.node_parser import SentenceSplitter
    from llama_index.embeddings.huggingface import HuggingFaceEmbedding
    from llama_index.vector_stores.chroma import ChromaVectorStore
    Settings.embed_model = HuggingFaceEmbedding(
        model_name=EMBED_MODEL_NAME,
        cache_folder=os.getenv("HF_HOME", str(Path.home() / ".cache/hf")),
        query_instruction="query: ",
        text_instruction="passage: ",
    )
    Settings.llm = None
    LLAMA_OK = True
except ImportError:
    LLAMA_OK = False

# ── PyMuPDF para PDFs (lectura robusta) ───────────────────────────────────────
try:
    import fitz  # PyMuPDF
    FITZ_OK = True
except ImportError:
    FITZ_OK = False

# ── python-docx para Word ─────────────────────────────────────────────────────
try:
    import docx as docxlib  # python-docx
    DOCX_OK = True
except ImportError:
    DOCX_OK = False


# ══════════════════════════════════════════════════════════════════════════════
# CONFIGURACIÓN DE CHUNKING
# ══════════════════════════════════════════════════════════════════════════════
# e5-small soporta hasta 512 tokens. Usamos chunks de 200 tokens con overlap
# de 40. Esto da chunks lo suficientemente grandes para tener contexto rico
# pero sin sobrepasar el límite del modelo.
CHUNK_SIZE_TOKENS    = int(os.getenv("CHUNK_SIZE_TOKENS",    "200"))
CHUNK_OVERLAP_TOKENS = int(os.getenv("CHUNK_OVERLAP_TOKENS", "40"))


# ══════════════════════════════════════════════════════════════════════════════
# HELPERS ChromaDB
# ══════════════════════════════════════════════════════════════════════════════

def _admin():
    return chromadb.AdminClient(settings=chromadb.Settings(
        chroma_api_impl="chromadb.api.fastapi.FastAPI",
        chroma_server_host=CHROMA_HOST,
        chroma_server_http_port=CHROMA_PORT,
    ))

def _client(tenant: str):
    return chromadb.HttpClient(
        host=CHROMA_HOST, port=CHROMA_PORT,
        tenant=tenant, database=CHROMA_DATABASE,
        settings=chromadb.Settings(anonymized_telemetry=False),
    )

def _docs(tenant: str) -> Path:
    p = DOCS_BASE_PATH / tenant
    p.mkdir(parents=True, exist_ok=True)
    return p

def _ensure(tenant: str):
    admin = _admin()
    try: admin.get_tenant(tenant)
    except: admin.create_tenant(tenant)
    try: admin.get_database(CHROMA_DATABASE, tenant=tenant)
    except: admin.create_database(CHROMA_DATABASE, tenant=tenant)

def _load_tenants() -> list[str]:
    DOCS_BASE_PATH.mkdir(parents=True, exist_ok=True)
    if not TENANTS_FILE.exists():
        _save_tenants([])
    try:
        return json.loads(TENANTS_FILE.read_text())
    except Exception:
        return []

def _save_tenants(tenants: list[str]):
    DOCS_BASE_PATH.mkdir(parents=True, exist_ok=True)
    TENANTS_FILE.write_text(json.dumps(tenants))


# ══════════════════════════════════════════════════════════════════════════════
# DETECTOR DE BOILERPLATE (cabeceras/pies de página repetidos)
# ══════════════════════════════════════════════════════════════════════════════

def _detect_boilerplate(pages_text: list[str], min_repeats: int = 3) -> list[str]:
    """Detecta líneas que se repiten en muchas páginas (cabeceras, pies)."""
    line_counter: dict = defaultdict(int)
    for page in pages_text:
        for line in {ln.strip() for ln in page.splitlines() if len(ln.strip()) >= 8}:
            line_counter[line] += 1

    threshold = max(min_repeats, len(pages_text) // 3)
    return [line for line, count in line_counter.items() if count >= threshold]


_GENERIC_BOILERPLATE_PATTERNS = [
    re.compile(r"^Page\s+\d+\s+of\s+\d+\s*$", re.IGNORECASE),
    re.compile(r"^\s*\d+\s*$"),
    re.compile(r"^Tel\s*:", re.IGNORECASE),
    re.compile(r"^Fax\s*:", re.IGNORECASE),
    re.compile(r"^www\.", re.IGNORECASE),
    re.compile(r"^\S+@\S+\.\S+\s*$"),
]


def _strip_boilerplate(text: str, boilerplate: list[str]) -> str:
    if not text:
        return text
    lines_to_drop = set(boilerplate)
    kept = []
    for line in text.splitlines():
        stripped = line.strip()
        if not stripped:
            kept.append(line)
            continue
        if stripped in lines_to_drop:
            continue
        if any(p.match(stripped) for p in _GENERIC_BOILERPLATE_PATTERNS):
            continue
        kept.append(line)
    return "\n".join(kept)


# ══════════════════════════════════════════════════════════════════════════════
# LECTORES POR FORMATO
# ══════════════════════════════════════════════════════════════════════════════

def _load_pdf(file_path: str) -> list:
    """Extrae texto página a página con PyMuPDF y elimina boilerplate."""
    if not FITZ_OK:
        raise RuntimeError("PyMuPDF no instalado. Ejecuta: pip install pymupdf")

    docs: list = []
    fname = Path(file_path).name
    try:
        pdf = fitz.open(file_path)
    except Exception as e:
        print(f"[RAG] Error abriendo {fname}: {e}")
        return docs

    total_pages = pdf.page_count

    pages_raw: list = []
    for page_num, page in enumerate(pdf, start=1):
        text = (page.get_text("text") or "").strip()
        pages_raw.append((page_num, text))
    pdf.close()

    boilerplate = _detect_boilerplate([t for _, t in pages_raw])
    if boilerplate:
        print(f"[RAG] {fname}: {len(boilerplate)} líneas-boilerplate detectadas, eliminadas.")

    skipped = []
    for page_num, raw in pages_raw:
        cleaned = _strip_boilerplate(raw, boilerplate)
        cleaned = re.sub(r"\n{3,}", "\n\n", cleaned).strip()

        alpha_chars = sum(1 for c in cleaned if c.isalpha())
        if alpha_chars < 30:
            skipped.append(page_num)
            continue

        docs.append(Document(
            text=cleaned,
            metadata={
                "file_name":   fname,
                "page_label":  str(page_num),
                "total_pages": total_pages,
                "source_type": "pdf",
            },
        ))

    if skipped:
        print(f"[RAG] {fname}: {len(docs)}/{total_pages} páginas con texto útil. "
              f"Descartadas: {skipped[:10]}{'...' if len(skipped) > 10 else ''}")
    return docs


def _load_docx(file_path: str) -> list:
    if not DOCX_OK:
        raise RuntimeError("python-docx no instalado. Ejecuta: pip install python-docx")

    fname = Path(file_path).name
    try:
        doc = docxlib.Document(file_path)
    except Exception as e:
        print(f"[RAG] Error abriendo {fname}: {e}")
        return []

    parts = []
    for para in doc.paragraphs:
        if para.text.strip():
            parts.append(para.text.strip())
    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(c.text.strip() for c in row.cells if c.text.strip())
            if row_text:
                parts.append(row_text)

    full_text = "\n\n".join(parts).strip()
    if len(full_text) < 30:
        return []

    return [Document(
        text=full_text,
        metadata={
            "file_name":   fname,
            "page_label":  "1",
            "source_type": "docx",
        },
    )]


def _load_textfile(file_path: str) -> list:
    fname = Path(file_path).name
    try:
        text = Path(file_path).read_text(encoding="utf-8", errors="replace").strip()
    except Exception as e:
        print(f"[RAG] Error leyendo {fname}: {e}")
        return []
    if len(text) < 30:
        return []
    ext = Path(file_path).suffix.lower()
    return [Document(
        text=text,
        metadata={
            "file_name":   fname,
            "page_label":  "1",
            "source_type": ext.lstrip("."),
        },
    )]


def _load_any(file_path: str) -> list:
    ext = Path(file_path).suffix.lower()
    if ext == ".pdf":
        return _load_pdf(file_path)
    if ext == ".docx":
        return _load_docx(file_path)
    if ext in (".txt", ".md"):
        return _load_textfile(file_path)
    raise ValueError(f"Formato no soportado: {ext}. Soportados: {sorted(SUPPORTED_EXTS)}")


# ══════════════════════════════════════════════════════════════════════════════
# APP
# ══════════════════════════════════════════════════════════════════════════════

app = FastAPI(title="RAG Manager API", version="1.2.0")
app.add_middleware(CORSMiddleware, allow_origins=["*"], allow_methods=["*"], allow_headers=["*"])

class TenantBody(BaseModel):
    name: str

class VectorizeBody(BaseModel):
    filenames: Optional[list[str]] = None


@app.get("/tenants")
def list_tenants():
    return {"tenants": _load_tenants()}

@app.post("/tenants")
def create_tenant(body: TenantBody):
    name = body.name.strip().lower().replace(" ", "_")
    if not name:
        raise HTTPException(400, "Nombre vacío")
    try:
        _ensure(name)
        tenants = _load_tenants()
        if name not in tenants:
            tenants.append(name)
            _save_tenants(tenants)
        return {"status": "ok", "tenant": name}
    except Exception as e:
        raise HTTPException(500, str(e))

@app.get("/tenants/{tenant}/documents")
def get_documents(tenant: str):
    try:
        _ensure(tenant)
        db  = _client(tenant)
        col = db.get_or_create_collection(COLLECTION_NAME)
        total   = col.count()
        indexed = {}
        if total > 0:
            result   = col.get(include=["metadatas"])
            docs_map = defaultdict(lambda: {"chunks": 0, "pages": set()})
            for meta in result["metadatas"]:
                fname = meta.get("file_name") or meta.get("filename") or "sin_nombre"
                docs_map[fname]["chunks"] += 1
                docs_map[fname]["pages"].add(str(meta.get("page_label", "?")))
            indexed = {n: {"chunks": d["chunks"], "pages": len(d["pages"])} for n, d in docs_map.items()}
        disk_files = []
        for ext in SUPPORTED_EXTS:
            disk_files.extend(f.name for f in sorted(_docs(tenant).glob(f"*{ext}")))
        disk_files = sorted(disk_files)
        return {
            "tenant": tenant, "total_chunks": total,
            "indexed": indexed, "disk_files": disk_files,
            "not_indexed": [f for f in disk_files if f not in indexed],
            "supported_extensions": sorted(SUPPORTED_EXTS),
        }
    except Exception as e:
        raise HTTPException(500, str(e))

@app.post("/tenants/{tenant}/upload")
async def upload(tenant: str, file: UploadFile = File(...)):
    ext = Path(file.filename).suffix.lower()
    if ext not in SUPPORTED_EXTS:
        raise HTTPException(400, f"Formato no soportado: {ext}. Soportados: {sorted(SUPPORTED_EXTS)}")
    try:
        _ensure(tenant)
        dest = _docs(tenant) / file.filename
        dest.write_bytes(await file.read())
        return {"status": "ok", "file": file.filename, "format": ext.lstrip(".")}
    except Exception as e:
        raise HTTPException(500, str(e))

@app.post("/tenants/{tenant}/vectorize")
def vectorize(tenant: str, body: VectorizeBody = VectorizeBody()):
    if not LLAMA_OK:
        raise HTTPException(503, "LlamaIndex no disponible")
    try:
        _ensure(tenant)
        docs_path = _docs(tenant)

        if body.filenames:
            to_process = [str(docs_path / f) for f in body.filenames
                          if (docs_path / f).exists()
                          and Path(f).suffix.lower() in SUPPORTED_EXTS]
        else:
            db  = _client(tenant)
            col = db.get_or_create_collection(COLLECTION_NAME)
            already = (
                {m.get("file_name", "") for m in col.get(include=["metadatas"])["metadatas"]}
                if col.count() > 0 else set()
            )
            to_process = []
            for ext in SUPPORTED_EXTS:
                to_process.extend(
                    str(f) for f in docs_path.glob(f"*{ext}")
                    if f.name not in already
                )

        if not to_process:
            return {
                "status": "ok",
                "processed": 0,
                "chunks_total": _client(tenant).get_or_create_collection(COLLECTION_NAME).count(),
            }

        all_docs: list = []
        per_file_pages: dict = {}
        errors: list = []
        for fpath in to_process:
            try:
                file_docs = _load_any(fpath)
                per_file_pages[Path(fpath).name] = len(file_docs)
                all_docs.extend(file_docs)
            except Exception as e:
                err_msg = f"{Path(fpath).name}: {e}"
                errors.append(err_msg)
                print(f"[RAG] {err_msg}")

        if not all_docs:
            return {
                "status": "ok",
                "processed": 0,
                "warning": "Ningún fichero produjo texto extraíble. "
                           "Si son PDFs escaneados, necesitan OCR (no incluido).",
                "per_file_pages": per_file_pages,
                "errors": errors,
            }

        splitter = SentenceSplitter(
            chunk_size=CHUNK_SIZE_TOKENS,
            chunk_overlap=CHUNK_OVERLAP_TOKENS,
        )

        db  = _client(tenant)
        col = db.get_or_create_collection(COLLECTION_NAME)
        VectorStoreIndex.from_documents(
            all_docs,
            storage_context=StorageContext.from_defaults(
                vector_store=ChromaVectorStore(chroma_collection=col)
            ),
            transformations=[splitter],
            show_progress=False,
        )

        return {
            "status": "ok",
            "processed": len(per_file_pages),
            "files": list(per_file_pages.keys()),
            "per_file_pages": per_file_pages,
            "chunks_total": col.count(),
            "embed_model":  EMBED_MODEL_NAME,
            "chunk_config": {
                "chunk_size_tokens":    CHUNK_SIZE_TOKENS,
                "chunk_overlap_tokens": CHUNK_OVERLAP_TOKENS,
            },
            "errors": errors,
        }
    except Exception as e:
        raise HTTPException(500, str(e))

@app.delete("/tenants/{tenant}/documents/{filename}")
def delete_document(tenant: str, filename: str, delete_file: bool = False):
    try:
        db  = _client(tenant)
        col = db.get_or_create_collection(COLLECTION_NAME)
        result = col.get(include=["metadatas"])
        ids = [result["ids"][i] for i, m in enumerate(result["metadatas"])
               if (m.get("file_name") or m.get("filename") or "") == filename]

        deleted_chunks = 0
        if ids:
            for i in range(0, len(ids), 100):
                col.delete(ids=ids[i:i+100])
            deleted_chunks = len(ids)
        elif not delete_file:
            raise HTTPException(404, f"'{filename}' no encontrado en el índice")

        file_deleted = False
        if delete_file:
            fp = _docs(tenant) / filename
            if fp.exists():
                fp.unlink()
                file_deleted = True

        return {"status": "ok", "deleted_chunks": deleted_chunks, "file_deleted": file_deleted}
    except HTTPException: raise
    except Exception as e: raise HTTPException(500, str(e))

@app.delete("/tenants/{tenant}/clear")
def clear_index(tenant: str):
    try:
        db = _client(tenant)
        db.delete_collection(COLLECTION_NAME)
        db.get_or_create_collection(COLLECTION_NAME)
        return {"status": "ok"}
    except Exception as e:
        raise HTTPException(500, str(e))

@app.delete("/tenants/{tenant}")
def delete_tenant(tenant: str):
    if tenant == "default_tenant":
        raise HTTPException(400, "No se puede eliminar el tenant por defecto")
    try:
        try:
            db = _client(tenant)
            db.delete_collection(COLLECTION_NAME)
        except Exception:
            pass
        docs_path = _docs(tenant)
        if docs_path.exists():
            shutil.rmtree(docs_path)
        tenants = _load_tenants()
        tenants = [t for t in tenants if t != tenant]
        _save_tenants(tenants)
        return {"status": "ok", "tenant": tenant}
    except HTTPException: raise
    except Exception as e: raise HTTPException(500, str(e))

@app.get("/health")
def health():
    try:
        httpx.get(f"http://{CHROMA_HOST}:{CHROMA_PORT}/api/v2/heartbeat", timeout=3)
        return {
            "status": "ok",
            "chroma": f"{CHROMA_HOST}:{CHROMA_PORT}",
            "embed_model": EMBED_MODEL_NAME,
            "supported_formats": sorted(SUPPORTED_EXTS),
            "readers": {
                "pdf":  "PyMuPDF" if FITZ_OK else "NO DISPONIBLE — pip install pymupdf",
                "docx": "python-docx" if DOCX_OK else "NO DISPONIBLE — pip install python-docx",
                "txt":  "stdlib",
                "md":   "stdlib",
            },
            "chunk_config": {
                "chunk_size_tokens":    CHUNK_SIZE_TOKENS,
                "chunk_overlap_tokens": CHUNK_OVERLAP_TOKENS,
            },
        }
    except:
        return {"status": "error", "chroma": "no disponible"}


@app.post("/upload")
async def upload_flat(
    file: UploadFile = File(...),
    tenant: str = Form("default_tenant"),
    collection: str = Form("documents"),
    chroma_db: str = Form("default_database"),
    filename: str = Form(""),
):
    if filename:
        file.filename = filename
    return await upload(tenant=tenant, file=file)


if __name__ == "__main__":
    import uvicorn
    uvicorn.run(app, host="0.0.0.0", port=int(os.getenv("PORT", "8300")))